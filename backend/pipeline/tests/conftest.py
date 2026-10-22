import io
import random

import httpx
import pytest

from kompas import curated, db
from kompas.http import Http
from kompas.periods import available_at, months


@pytest.fixture
def cur():
    return curated.load()


@pytest.fixture
def conn(tmp_path, monkeypatch):
    monkeypatch.setenv("KOMPAS_DB", str(tmp_path / "test.db"))
    monkeypatch.setenv("KOMPAS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    c = db.connect()
    yield c
    c.close()


def add_hiring(conn, period: str, n_total: int, n_matching: int, phrase: str):
    """Synthetic test data: n_total hiring posts in a month, n_matching of them mentioning phrase."""
    for i in range(n_total):
        text = f"Backend engineer, remote. We use {phrase} daily." if i < n_matching else "Backend engineer, Go."
        db.upsert_doc(conn, {
            "id": f"hnh:{period}:{i}", "source": "hn_hiring", "kind": "hiring_post",
            "url": "https://news.ycombinator.com/", "title": text[:40], "text": text,
            "published_at": f"{period}-02T12:00:00Z", "period": period, "metrics": {},
        })


def add_github_counts(conn, cur, period: str, topic_counts: dict[str, int], default: int = 100):
    for t in cur.all_topics():
        db.upsert_doc(conn, {
            "id": f"ghc:{t}:{period}", "source": "github", "kind": "repo_count", "url": "https://github.com/search",
            "title": f"New repositories tagged '{t}' in {period}", "text": "",
            "published_at": available_at(period), "period": period,
            "metrics": {"topic": t, "count": topic_counts.get(t, default)},
        })


def add_dosm(conn, period: str, release_date: str, rows: dict[str, int]):
    for code, n in rows.items():
        db.upsert_doc(conn, {
            "id": f"dosm:{period}:{code}", "source": "dosm_jmi", "kind": "occupation_stat",
            "url": "https://www.dosm.gov.my/", "title": f"Occupation {code}", "text": "",
            "published_at": f"{release_date}T00:00:00Z", "period": period,
            "metrics": {"masco_code": code, "offerings": n},
        })


def mock_http(handler, **kwargs) -> Http:
    """An Http client whose requests go to handler(request) -> httpx.Response. Never sleeps."""
    kwargs.setdefault("sleep", lambda s: None)
    kwargs.setdefault("rng", random.Random(0))
    return Http(transport=httpx.MockTransport(handler), **kwargs)


def routes(table: dict):
    """handler for mock_http: maps a URL path (or a callable on the request) to a response or a list of
    responses served in turn. Unknown paths answer 404."""
    served = {k: (list(v) if isinstance(v, list) else v) for k, v in table.items()}

    def handler(request: httpx.Request) -> httpx.Response:
        for key, value in served.items():
            if (callable(key) and key(request)) or key == request.url.path:
                if isinstance(value, list):
                    return value.pop(0) if len(value) > 1 else value[0]
                return value(request) if callable(value) else value
        return httpx.Response(404, text="not found")

    return handler


def jmi_workbook(values: dict[str, dict[str, object]], title_suffix: str = "by Major Group Occupation") -> bytes:
    """A workbook laid out like DOSM's JMI release file: Table 1 (economic activity) first, then the
    major-group table with annual, quarterly and monthly columns. values: {row label: {quarter: value}}."""
    import openpyxl

    quarters = sorted({q for row in values.values() for q in row})
    wb = openpyxl.Workbook()
    ws = wb.active
    years = sorted({int(q[:4]) for q in quarters})

    def table(top: int, title: str, label: str, rows: dict):
        ws.cell(top + 1, 2, title)
        for j, y in enumerate(years):                         # annual columns
            ws.cell(top + 3, 3 + j, y)
        col, seen = 11, set()
        for q in quarters:                                    # quarterly: year only on a year's first column
            if q[:4] not in seen:
                ws.cell(top + 3, col, int(q[:4]))
                seen.add(q[:4])
            ws.cell(top + 4, col, q[4:])
            col += 1
        ws.cell(top + 3, col + 1, years[0])                   # a monthly block after the quarters
        ws.cell(top + 4, col + 1, "January")
        ws.cell(top + 4, 2, label)
        r = top + 5
        for name, by_q in rows.items():
            ws.cell(r, 2, name)
            for j, q in enumerate(quarters):
                ws.cell(r, 11 + j, by_q.get(q))
            ws.cell(r, col + 1, 123)
            r += 1
        ws.cell(r, 2, "Source: Big Data Analytics: Job Market Insights, DOSM")
        return r + 2

    nxt = table(0, "Table 1: Number of Job Offering Advertised Online by Economic Activity", "Economic Activity",
                {"Jumlah": {q: 1 for q in quarters}, "Manufacturing": {q: 1 for q in quarters}})
    table(nxt, f"Table 2: Number of Job Offerings Advertised Online {title_suffix}", "Occupation", values)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["add_hiring", "add_github_counts", "add_dosm", "months", "mock_http", "routes", "jmi_workbook"]
