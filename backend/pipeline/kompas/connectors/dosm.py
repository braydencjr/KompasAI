"""DOSM Big Data Analytics: Job Market Insights (JMI), Malaysia's official source for online job offerings.

Two streams, because DOSM publishes two levels of detail in different ways:

groups (automated)
    Recent quarterly release pages on dosm.gov.my have an XLSX workbook attached. Its "by (Major Group)
    Occupation" table gives job offerings per MASCO 2020 *major group* (1 digit, e.g. "Professionals"), per
    quarter since 2019Q3. Releases are discovered from the official release archive and imported oldest
    first, so each quarter keeps the value and date of the release that first published it (a release
    only counts for quarters up to its own: workbooks can hold placeholder zeros for later ones); later
    revisions are kept as observations (artifact_metric 'offerings', observed at the revising release).
    Kind: occupation_group_stat. Not used for skill scoring: a major group such as "Professionals" spans
    accountants and software developers alike, so it can't stand in for any one skill.

occupations (manual)
    4-digit MASCO occupations are what the skill mapping needs (pipeline/curated/occupation_code_map.csv).
    DOSM shows them only in its interactive portal and supplies them on request; they aren't in the release
    workbooks. A CSV built from official figures is imported from data/manual/dosm_jmi.csv after strict
    validation. Kind: occupation_stat, which aggregate.py turns into rung-3 demand.

Release dates are known to the day, not the hour, so a release counts as public at the end of its
release day in Malaysian time (document.release_day_end): later rather than earlier.
"""

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from .. import config
from ..curated import Curated
from ..document import Document, release_day_end
from ..http import FetchError, Http
from ..ingest import DatasetInvalid, Options, Run
from ..periods import period_end

# ---------- manual CSV: 4-digit occupations ----------

COLUMNS = ("period", "release_date", "masco_code", "occupation", "offerings", "source_url")
MISSING = {"", "-", "na", "n/a", "n.a.", "..", "x"}


@dataclass
class CsvRow:
    line: int
    period: str
    release_date: str
    masco_code: str
    occupation: str
    offerings: int
    source_url: str


def _quarters(first: str, last: str) -> list[str]:
    y, q = int(first[:4]), int(first[5])
    out = []
    while f"{y}Q{q}" <= last:
        out.append(f"{y}Q{q}")
        y, q = (y + 1, 1) if q == 4 else (y, q + 1)
    return out


def validate_csv(text: str, cur: Curated) -> tuple[list[CsvRow], list[str], list[str]]:
    """(rows, errors, warnings). Any error rejects the whole file; warnings are recorded and the file is
    imported. A missing value is never read as zero: its row is left out and reported."""
    errors, warnings, rows = [], [], []
    reader = csv.DictReader(io.StringIO(text))
    missing = set(COLUMNS) - set(reader.fieldnames or ())
    if missing:
        return [], [f"missing columns: {', '.join(sorted(missing))}"], []
    extra = set(reader.fieldnames or ()) - set(COLUMNS)
    if extra:
        warnings.append(f"extra columns ignored: {', '.join(sorted(extra))}")
    mapped = {code for (system, code) in cur.occupation_codes if system == "MASCO"}
    seen: dict[tuple[str, str], int] = {}
    release_of: dict[str, str] = {}
    for line, raw in enumerate(reader, start=2):
        r = {k: (v or "").strip() for k, v in raw.items() if k}
        if not any(r.values()) or r["period"].startswith("#"):
            continue
        where = f"line {line}"
        bad = []
        if not re.fullmatch(r"\d{4}Q[1-4]", r["period"]):
            bad.append(f"period {r['period']!r} is not YYYYQn")
        try:
            released = date.fromisoformat(r["release_date"])
        except ValueError:
            released = None
            bad.append(f"release_date {r['release_date']!r} is not YYYY-MM-DD")
        if released and not bad and released <= period_end(r["period"]):
            bad.append(f"release_date {released} is not after the end of {r['period']}")
        if not re.fullmatch(r"\d{4}", r["masco_code"]):
            bad.append(f"masco_code {r['masco_code']!r} is not a 4-digit MASCO code")
        if not r["occupation"]:
            bad.append("occupation is empty")
        value = r["offerings"].replace(",", "").replace(" ", "")
        offerings = None
        if value.lower() in MISSING:
            warnings.append(f"{where}: no offerings figure for {r['period']} {r['masco_code']}; row left out "
                            "(a missing figure is not zero)")
        elif not re.fullmatch(r"\d+", value):
            bad.append(f"offerings {r['offerings']!r} is not a whole non-negative number")
        else:
            offerings = int(value)
        if not r["source_url"].startswith("https://"):
            bad.append(f"source_url {r['source_url']!r} is not an https URL")
        elif not re.match(r"https://([\w-]+\.)*dosm\.gov\.my/", r["source_url"]):
            warnings.append(f"{where}: source_url is not on dosm.gov.my: {r['source_url']}")
        if bad:
            errors.extend(f"{where}: {b}" for b in bad)
            continue
        key = (r["period"], r["masco_code"])
        if key in seen:
            errors.append(f"{where}: duplicate {r['period']} {r['masco_code']} (also on line {seen[key]})")
            continue
        seen[key] = line
        if release_of.setdefault(r["period"], r["release_date"]) != r["release_date"]:
            errors.append(f"{where}: {r['period']} has two release dates ({release_of[r['period']]} and "
                          f"{r['release_date']}); one release publishes a quarter")
            continue
        if r["masco_code"] not in mapped:
            warnings.append(f"{where}: MASCO {r['masco_code']} isn't in occupation_code_map.csv; it counts "
                            "toward the quarter's total but no skill")
        if offerings is not None:
            rows.append(CsvRow(line, r["period"], r["release_date"], r["masco_code"], r["occupation"], offerings,
                               r["source_url"]))
    if not rows and not errors:
        errors.append("no data rows")
    periods = sorted({r.period for r in rows})
    if periods:
        gaps = [p for p in _quarters(periods[0], periods[-1]) if p not in periods]
        if gaps:
            warnings.append(f"no rows for {', '.join(gaps)}: trends across these quarters can't be computed")
        codes = {p: {r.masco_code for r in rows if r.period == p} for p in periods}
        for a, b in zip(periods, periods[1:]):
            if codes[a] != codes[b]:
                warnings.append(f"the occupations listed change between {a} and {b}; shares are relative to the "
                                "listed occupations, so this shifts every share")
    return rows, errors, warnings


def ingest_csv(run: Run, cur: Curated, opts: Options):
    """Imports data/manual/dosm_jmi.csv (or opts.file) after validation. Nothing is written unless the
    whole file validates. Re-importing the same file is a no-op; a changed figure for a quarter already
    stored is kept as a revision observation, and the first-published figure stays in the document."""
    path = Path(opts.file or config.MANUAL_DIR / "dosm_jmi.csv")
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    run.detail.update(file=str(path), sha256=digest)
    if run.checkpoint() == digest and not opts.force:
        run.skip()
        return
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise DatasetInvalid([f"{path.name} is not UTF-8 text: {e}"]) from e
    rows, errors, warnings = validate_csv(text, cur)
    for w in warnings:
        run.warn(path.name, "validation", w)
    if errors:
        raise DatasetInvalid(errors)
    for r in rows:
        doc_id = f"dosm:{r.period}:{r.masco_code}"
        published = release_day_end(r.release_date)
        existing = run.conn.execute("SELECT metrics FROM raw_doc WHERE id = ?", (doc_id,)).fetchone()
        old = json.loads(existing["metrics"]).get("offerings") if existing else None
        if existing and old != r.offerings:
            run.warn(doc_id, "revision", f"{r.period} {r.masco_code}: {old} -> {r.offerings}; the first-published "
                     "figure stays in use and the new one is kept as an observation")
            run.skip()
        else:
            run.save(Document(
                id=doc_id, source="dosm_jmi", kind="occupation_stat", source_id=f"{r.period}:{r.masco_code}",
                url=r.source_url, title=r.occupation, text="", published_at=published, period=r.period,
                metrics={"masco_code": r.masco_code, "offerings": r.offerings},
                raw_payload={"file": path.name, "sha256": digest, "line": r.line, "period": r.period,
                             "release_date": r.release_date, "masco_code": r.masco_code, "occupation": r.occupation,
                             "offerings": r.offerings, "source_url": r.source_url},
            ))
        run.observe(doc_id, "offerings", r.offerings, published)
    periods = sorted({r.period for r in rows})
    run.detail.update(rows=len(rows), periods=periods)
    run.commit(cursor=digest)


# ---------- automated: official release workbooks, major groups ----------

SERIES = re.compile(r"job market insights|job vacancies", re.I)
ORDINAL = {"first": 1, "second": 2, "third": 3, "fourth": 4}

# MASCO 2020 major groups, matched on the start of the row label (labels vary a little between releases)
MAJOR_GROUPS = (
    ("jumlah", "TOTAL"), ("total", "TOTAL"), ("managers", "1"), ("professionals", "2"), ("technicians", "3"),
    ("clerical", "4"), ("service and sales", "5"), ("skilled agricultural", "6"), ("craft", "7"),
    ("plant and machine", "8"), ("elementary", "9"), ("unclassified", "X"), ("not classified", "X"),
)
GROUP_TITLES = {"1": "Managers", "2": "Professionals", "3": "Technicians and associate professionals",
                "4": "Clerical support workers", "5": "Service and sales workers",
                "6": "Skilled agricultural, forestry, livestock and fishery workers",
                "7": "Craft and related trades workers", "8": "Plant and machine operators and assemblers",
                "9": "Elementary occupations", "X": "Unclassified", "TOTAL": "All occupations"}


class TableError(ValueError):
    """The workbook doesn't have the expected table layout."""


@dataclass(frozen=True)
class Release:
    slug: str
    date: str          # YYYY-MM-DD
    title: str
    quarter: str | None

    @property
    def url(self) -> str:
        return f"{config.DOSM_JMI_BASE}/portal-main/release-content/{self.slug}"


def parse_archive(html_text: str) -> list[Release]:
    out = []
    for slug, body in re.findall(r'href="/portal-main/release-content/([^"]+)"[^>]*>(.*?)</a>', html_text, re.S):
        label = " ".join(re.sub(r"<[^>]+>", " ", body).split())
        m = re.match(r"(\d\d)\.(\d\d)\.(\d{4})\s*-\s*(.+)", label)
        if not m or not SERIES.search(m.group(4)):
            continue
        q = re.search(r"(first|second|third|fourth) quarter of (\d{4})", m.group(4), re.I)
        out.append(Release(slug, f"{m.group(3)}-{m.group(2)}-{m.group(1)}", m.group(4),
                           f"{q.group(2)}Q{ORDINAL[q.group(1).lower()]}" if q else None))
    return out


def discover(http: Http) -> list[Release]:
    """Every JMI release listed in the official archive, oldest first."""
    found: dict[str, Release] = {}
    last_error = None
    for anchor in config.DOSM_JMI_ANCHORS:
        try:
            page = http.get(f"{config.DOSM_JMI_BASE}/portal-main/release-archive/{anchor}").text
        except FetchError as e:
            last_error = e
            continue
        for r in parse_archive(page):
            found.setdefault(r.slug, r)
    if not found:
        raise last_error or FetchError("parse", "no Job Market Insights releases found in the DOSM archive")
    return sorted(found.values(), key=lambda r: (r.date, r.slug))


def document_links(html_text: str) -> list[str]:
    """Attachments on a release page that aren't marked as PDFs (the data workbook is among these)."""
    links = []
    for m in re.finditer(r"release-document-log\?release_document_id=(\d+)([^\"'\s<]*)", html_text):
        if "type=pdf" not in m.group(2):
            url = f"{config.DOSM_JMI_BASE}/portal-main/release-document-log?release_document_id={m.group(1)}"
            if url not in links:
                links.append(url)
    return links


def upcoming_release(html_text: str) -> str | None:
    m = re.search(r"Upcoming release\s*:\s*(\d{1,2} \w+ \d{4})", html_text)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%d %B %Y").date().isoformat()
    except ValueError:
        return None


def _group(label: str) -> str | None:
    s = " ".join(label.lower().split())
    return next((code for prefix, code in MAJOR_GROUPS if s.startswith(prefix)), None)


def _number(v) -> int | None:
    """A cell's value: int, or None when DOSM shows it as not available ('-', blank)."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise TableError(f"unexpected value {v!r}")
    if isinstance(v, (int, float)):
        if v < 0 or v != int(v):
            raise TableError(f"offerings must be a whole non-negative number, got {v!r}")
        return int(v)
    s = str(v).strip().replace(",", "")
    if s.lower() in MISSING:
        return None
    if not s.isdigit():
        raise TableError(f"unexpected value {v!r}")
    return int(s)


def parse_groups_xlsx(data: bytes) -> dict[tuple[str, str], int | None]:
    """{(quarter, major group): offerings or None} from a JMI release workbook's major-group table."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:   # openpyxl raises several unrelated types for a corrupt file
        raise TableError(f"not a readable XLSX workbook: {type(e).__name__}") from e
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        title_at = next((i for i, r in enumerate(rows)
                         if any(isinstance(c, str) and re.search(r"by\s+(major\s+group\s+)?occupation\s*$", c.strip(), re.I) for c in r)),
                        None)
        if title_at is None:
            continue
        label_at = next((i for i in range(title_at + 1, min(title_at + 8, len(rows)))
                         if any(isinstance(c, str) and c.strip().lower() == "occupation" for c in rows[i])), None)
        if label_at is None:
            raise TableError("major-group table has no 'Occupation' header row")
        labels, years = rows[label_at], next((rows[i] for i in range(label_at - 1, title_at, -1) if any(rows[i])), [])
        columns, year = {}, None
        for j, lab in enumerate(labels):
            y = years[j] if j < len(years) else None
            if isinstance(y, (int, float)) or (isinstance(y, str) and y.strip().isdigit()):
                year = int(y)
            if isinstance(lab, str) and re.fullmatch(r"Q[1-4]", lab.strip()):
                if year is None:
                    raise TableError(f"quarter column {j} has no year")
                period = f"{year}Q{lab.strip()[1]}"
                if period in columns.values():
                    raise TableError(f"{period} appears twice in the header")
                columns[j] = period
        if not columns:
            raise TableError("major-group table has no quarterly columns")
        ordered = list(columns.values())
        if ordered != sorted(ordered):
            raise TableError("quarterly columns are out of order")
        values, groups = {}, set()
        for r in rows[label_at + 1:]:
            label = next((c for c in r if c is not None), None)
            if label is None or (isinstance(label, str) and label.strip().lower().startswith("source")):
                break
            group = _group(str(label))
            if group is None:
                raise TableError(f"unknown occupation group {label!r}")
            if group in groups:
                raise TableError(f"occupation group {label!r} appears twice")
            groups.add(group)
            for j, period in columns.items():
                values[(period, group)] = _number(r[j] if j < len(r) else None)
        if "TOTAL" not in groups or not {"1", "2", "3", "4"} <= groups:
            raise TableError(f"major-group table is missing rows (found {sorted(groups)})")
        return values
    raise TableError("no job-offerings-by-occupation table in the workbook")


def check_totals(values: dict[tuple[str, str], int | None]) -> dict[str, str]:
    """{quarter: problem} for quarters whose reported groups don't add up to the published total (more
    than 0.5% apart). Groups shown as not available ('-', e.g. Unclassified) are left out of the sum."""
    out = {}
    for period in sorted({p for p, _ in values}):
        total = values.get((period, "TOTAL"))
        parts = [v for (p, g), v in values.items() if p == period and g != "TOTAL" and v is not None]
        if total is None or not parts:
            continue
        if abs(sum(parts) - total) > 0.005 * max(total, 1):
            out[period] = f"{period}: groups add up to {sum(parts):,}, published total is {total:,}"
    return out


def _workbook(http: Http, release: Release) -> tuple[bytes | None, str, list[str]]:
    """(workbook bytes or None, release page, notes). An attachment link DOSM answers with 4xx is dead for
    good (some 2022 releases), so it's noted and skipped; server and network errors still raise, because
    those may pass and later releases shouldn't be imported before this one."""
    page, notes = http.get(release.url).text, []
    for link in document_links(page):
        try:
            resp = http.get(link)
        except FetchError as e:
            if e.category != "client":
                raise
            notes.append(f"attachment {link.rsplit('=', 1)[-1]} answers HTTP {e.status}")
            continue
        kind = resp.headers.get("content-type", "")
        if "spreadsheetml" in kind or (resp.content[:4] == b"PK\x03\x04" and "pdf" not in kind):
            return resp.content, page, notes
    return None, page, notes


def ingest_groups(run: Run, cur: Curated, opts: Options, http: Http | None = None):
    """Imports every listed release's workbook not imported yet, oldest first. The checkpoint holds the
    imported releases and the ones without a usable workbook (retried only with --force)."""
    http = http or run.http(min_interval=1.0)
    state = json.loads(run.checkpoint() or "{}")
    imported, unusable = list(state.get("imported", [])), dict(state.get("unusable", {}))
    releases = discover(http)
    run.detail["releases_listed"] = len(releases)
    for rel in releases:
        if not opts.force and (rel.slug in imported or rel.slug in unusable):
            run.skip()
            continue
        try:
            data, page, notes = _workbook(http, rel)
            values = parse_groups_xlsx(data) if data else None
        except FetchError as e:
            # later releases wait, so each quarter keeps the date of the release that first published it
            run.fail(rel.slug, e.category, str(e))
            run.partial(f"{rel.slug} ({rel.date}) couldn't be fetched; later releases were not imported")
            break
        except TableError as e:
            unusable[rel.slug] = str(e)
            run.warn(rel.slug, "validation", f"workbook not usable: {e}")
            run.commit(cursor=_state(imported, unusable, state))
            continue
        if values is None:
            unusable[rel.slug] = "; ".join(notes) or "no XLSX workbook attached"
            run.warn(rel.slug, "validation", f"no usable XLSX workbook: {unusable[rel.slug]}")
            run.commit(cursor=_state(imported, unusable, state))
            continue
        inconsistent = check_totals(values)
        published = release_day_end(rel.date)
        for (period, group), value in sorted(values.items()):
            if value is None:
                continue   # "-" or blank: not available, which is not zero
            if rel.quarter and period > rel.quarter:
                continue   # a workbook can carry placeholder zeros for quarters it doesn't cover yet
            doc_id = f"dosmg:{period}:{group}"
            existing = run.conn.execute("SELECT metrics FROM raw_doc WHERE id = ?", (doc_id,)).fetchone()
            if existing is None and period in inconsistent:
                # DOSM has published placeholder zeros (2024Q3's total, 2025Q4's elementary workers) and
                # corrected them in the next release; a quarter that doesn't add up waits for that release
                run.skip()
            elif existing is None:
                run.save(Document(
                    id=doc_id, source="dosm_jmi", kind="occupation_group_stat", source_id=f"{period}:{group}",
                    url=rel.url, title=f"{GROUP_TITLES[group]}: online job offerings, {period}", text="",
                    published_at=published, period=period,
                    metrics={"masco_major_group": group, "offerings": value, "release": rel.slug,
                             "release_date": rel.date},
                    raw_payload={"release": rel.slug, "release_date": rel.date, "release_title": rel.title,
                                 "table": "Number of Job Offerings Advertised Online by Major Group Occupation",
                                 "period": period, "group": group, "offerings": value},
                ))
            else:
                old = json.loads(existing["metrics"])["offerings"]
                if old != value:
                    run.warn(doc_id, "revision", f"{rel.slug} revised {period} {GROUP_TITLES[group]}: {old:,} -> "
                             f"{value:,}; the first-published figure stays, the revision is kept as an observation")
                run.skip()
            run.observe(doc_id, "offerings", value, published)
        for period, problem in inconsistent.items():
            run.warn(rel.slug, "validation", problem + ("; not used as this quarter's first publication"
                                                         if not rel.quarter or period <= rel.quarter else ""))
        imported.append(rel.slug)
        state["latest_release"] = {"slug": rel.slug, "date": rel.date, "quarter": rel.quarter}
        state["upcoming_release"] = upcoming_release(page)
        run.commit(cursor=_state(imported, unusable, state))
    run.detail.update(imported=len(imported), unusable=len(unusable),
                      latest_release=state.get("latest_release"), upcoming_release=state.get("upcoming_release"))


def _state(imported: list[str], unusable: dict, state: dict) -> str:
    return json.dumps({**state, "imported": imported, "unusable": unusable}, sort_keys=True)
